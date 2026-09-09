# -*- coding: utf-8 -*-
"""Propuesta y presupuesto del Sistema de Ventas, en .docx."""

import datos as D
from comun import (doc, p, vineta, tabla, aviso, precio, salto, sin,
                   AZUL, GRIS)
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

U = lambda x: sin('US$ ', D.m(x))
B = lambda x: sin('Bs ', D.b(x))

# ===================================================================== PORTADA
for _ in range(5):
    doc.add_paragraph()

p('PROPUESTA Y PRESUPUESTO', negrita=True, tam=12, color=GRIS, centro=True, espacio=2)

_t = doc.add_paragraph()
_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
_r = _t.add_run('Sistema de Ventas')
_r.bold = True
_r.font.size = Pt(34)
_r.font.color.rgb = AZUL

p('Punto de venta, caja, almacén y comprobantes', tam=13.5, color=GRIS, centro=True, espacio=2)
p('para minimarkets y tiendas de barrio', tam=13.5, color=GRIS, centro=True, espacio=28)

p('Preparado para', negrita=True, tam=10, color=GRIS, centro=True, espacio=2)
p('Minimarket El Ahorro', tam=13, centro=True, espacio=20)

p('Santa Cruz de la Sierra, Bolivia', tam=10, color=GRIS, centro=True, espacio=2)
p('9 de septiembre de 2026 · válida por 30 días', tam=10, color=GRIS, centro=True, espacio=28)

p('Las cifras se expresan en dólares estadounidenses y su equivalente en bolivianos, '
  'al tipo de cambio de Bs 7,00 por US$.', tam=9.5, color=GRIS, centro=True, cursiva=True)

salto()

# ================================================================== 1. RESUMEN
doc.add_heading('1. En una página', level=1)

p('Un minimarket pierde plata en tres lugares a la vez, y casi siempre sin enterarse: en el '
  'mostrador, donde el precio depende de que el cajero se acuerde; en el cajón, donde al cerrar '
  'nadie sabe cuánto debería haber; y en el estante, donde el producto que más rota se termina '
  'un sábado a la tarde porque nadie llevaba la cuenta.')

p('Este sistema resuelve los tres. Y no es una promesa de desarrollo: está construido, probado '
  'y funcionando. Se instala y se empieza a usar.')

tabla(['', 'Resumen'],
      [['Qué es', 'Un sistema web de punto de venta con caja por turno, almacén, comprobantes, clientes, devoluciones y reportes.'],
       ['En qué estado está', 'Terminado. 275 pruebas automáticas, y una demo en línea con cuatro meses de operación ya cargada.'],
       ['Cómo se usa', 'Desde el navegador, en cualquier computadora del local. No hay que instalar nada en cada máquina.'],
       ['Qué se ofrece', 'Dos formas de acceder: comprarlo una sola vez, o pagar una suscripción por su uso.'],
       ['Cuánto cuesta', 'La compra, %s. La suscripción, %s al mes.'
        % (D.ambos(D.COMPRA_LANZAMIENTO_USD), D.ambos(D.SUSCRIPCION_MES_USD))]],
      anchos=[3.3, 12.2])

aviso('Se puede probar antes de decidir nada.',
      'La demo está publicada en internet con cuatro meses de ventas, caja, almacén y '
      'comprobantes cargados. Se entrega el enlace y tres cuentas de prueba —gerente, cajero y '
      'almacenero— para que el personal lo use unos días antes de que se firme nada.')

salto()

# ================================================================ 2. QUÉ HACE
doc.add_heading('2. Qué hace el sistema', level=1)

doc.add_heading('2.1 En el mostrador', level=2)
p('La pantalla de cobro está pensada para que una venta simple se complete sin soltar el lector '
  'de código de barras: se pasa el producto, Enter, y cobra.')
vineta('el precio sale siempre del catálogo. El cajero no puede escribirlo a mano.',
       'El precio no se decide en la caja: ')
vineta('por encima del máximo configurado hace falta la autorización de un supervisor.',
       'El descuento tiene techo: ')
vineta('lo calcula el sistema y lo muestra en pantalla. Se acaban las cuentas de cabeza.',
       'El vuelto no se calcula a mano: ')
vineta('una venta puede repartirse entre varias formas —una parte por QR, el resto en efectivo— '
       'y una de las líneas puede quedar en blanco para llevarse lo que falte.',
       'El pago se puede partir: ')
vineta('si no hay stock, la venta no se registra.', 'No se vende lo que no está: ')

doc.add_heading('2.2 Cobro por QR', level=2)
p('El sistema genera un código QR con el importe ya metido dentro: el cliente apunta el teléfono '
  'y paga exactamente lo que debe, sin teclear el monto y sin equivocarse.')
p('La venta se registra recién cuando el pago está confirmado. Es al revés de como suele hacerse, '
  'y es a propósito: una venta creada antes de cobrar ya descontó stock y ya emitió comprobante, '
  'así que si el cliente se arrepiente queda mercadería descontada que nadie se llevó.')
aviso('Requiere convenio con el banco.',
      'La parte del sistema está lista y probada. Falta que el negocio pida a su banco el acceso '
      'a la API de cobros por QR. Mientras tanto el QR se genera igual y el cajero confirma el '
      'pago viendo el comprobante en el teléfono del cliente, y queda registrado con su nombre. '
      'Conectar el banco no cuesta desarrollo adicional: se completan tres datos de '
      'configuración.')

doc.add_heading('2.3 La caja, por turno', level=2)
p('Cada cajero abre su turno con el efectivo con el que empieza, y lo cierra contando lo que hay. '
  'El sistema calcula cuánto debería haber y señala la diferencia. El cierre se imprime para que '
  'el cajero lo firme.')
p('Al cajón entra solo el efectivo. Lo cobrado con tarjeta, transferencia, billetera o QR se '
  'registra como venta, pero no se le reclama al cajero al cerrar: ese dinero está en el banco, '
  'no en su cajón. Lo mismo con las devoluciones, que salen prorrateadas.')

doc.add_heading('2.4 El almacén', level=2)
p('Existencias con alerta de lo que bajó del mínimo. Ingreso de mercadería a un producto que ya '
  'existe, sin darlo de alta otra vez. Ajustes por merma o rotura. Y un kardex que guarda cada '
  'movimiento con el stock que había antes y el que quedó después.')

doc.add_heading('2.5 Documentos y clientes', level=2)
p('El sistema elige solo el tipo de comprobante según el cliente: una persona recibe recibo, una '
  'empresa recibe factura. El cajero no tiene que acordarse. Se imprime en ticket de 80 mm o en '
  'hoja A4.')
p('Nada se borra nunca. Un documento se anula o se sustituye, y la cadena queda a la vista: quién '
  'lo hizo, cuándo y por qué.')

doc.add_heading('2.6 Devoluciones y reportes', level=2)
p('Devoluciones totales o parciales, decidiendo producto por producto si vuelve al estante o no: '
  'lo roto no vuelve. Y reportes de ventas por día, por forma de pago y por cajero, ranking de '
  'productos y alertas de reposición, exportables a Excel y a PDF.')

doc.add_heading('2.7 Quién puede hacer qué', level=2)
p('Tres roles de fábrica —gerente, cajero y almacenero— sobre una matriz de permisos que se puede '
  'ajustar. Cada cuenta entra directo a su pantalla de trabajo: el cajero al mostrador, quien '
  'lleva la gestión al panel.')

salto()

# ============================================================== 3. QUÉ SE ENTREGA
doc.add_heading('3. Qué se entrega', level=1)

tabla(['Entregable', 'Detalle'],
      [['El sistema funcionando', 'Instalado, con dominio, HTTPS y respaldos configurados.'],
       ['El catálogo cargado', 'Hasta 500 productos con su código, precio, unidad, categoría y stock inicial, a partir de la lista que entregue el negocio.'],
       ['Los datos del negocio', 'Nombre, dirección, teléfono e identificación fiscal, que salen impresos en cada comprobante.'],
       ['Las cuentas del personal', 'Una por persona, con su rol. Las de prueba se desactivan antes de entregar.'],
       ['Capacitación', 'Dos sesiones: una con quien cobra, otra con quien administra.'],
       ['Dos manuales', 'Uno de uso, con ejercicios paso a paso. Otro de funcionamiento, como referencia de cómo trabaja cada módulo.'],
       ['Soporte de arranque', 'Tres meses de correcciones sin costo. Un error se corrige; un cambio de alcance entra como orden de cambio.']],
      anchos=[4.0, 11.5])

doc.add_heading('3.1 Lo que NO incluye', level=2)
p('Se dice de frente, para que no haya sorpresas después.')
tabla(['Fuera de alcance', 'Por qué, y qué haría falta'],
      [['Facturación electrónica ante Impuestos Nacionales', 'El modelo de datos está preparado, pero la integración con el servicio del SIN no está en esta versión. Se cotiza aparte cuando el negocio esté obligado o quiera hacerlo.'],
       ['Módulo de compras a proveedores', 'Hoy la mercadería entra por «Ingreso» en el almacén, que resuelve el stock. Órdenes de compra, cuentas por pagar y costo promedio no están.'],
       ['Aplicación móvil', 'El sistema es web y se ve bien en el navegador del celular, pero no hay una app instalable.'],
       ['El equipamiento', 'Lector de código de barras, impresora térmica, computadora y conexión los pone el negocio.'],
       ['Servidor y dominio', 'Si el negocio no los tiene, se cotizan aparte. En la suscripción van incluidos.'],
       ['El convenio bancario para el QR', 'Lo gestiona el negocio con su banco. La parte del sistema ya está hecha.']],
      anchos=[4.8, 10.7])

salto()

# ============================================================ 4. CÓMO SE VALORÓ
doc.add_heading('4. Cómo se valoró el trabajo', level=1)

p('El precio no salió de una corazonada. Se calculó con el arancel del Colegio de Ingenieros en '
  'Tecnologías de la Información de Santa Cruz, y con las horas que el sistema tiene realmente '
  'encima, medidas contra el repositorio.')

doc.add_heading('4.1 La tarifa', level=2)
p('El arancel del CITI fija el valor de la hora de servicio profesional según la categoría del '
  'trabajo:')

tabla(['Categoría profesional', 'Bs / hora', 'US$ / hora', 'Cuándo aplica'],
      [[nom, B(t), U(t / D.TC), det] for nom, t, det in D.ARANCEL],
      anchos=[3.5, 1.9, 2.1, 8.0], destacar=(1,), derecha=(1, 2))

p('Se usa el escalón de técnico superior: corresponde a un desarrollo profesional a cargo del '
  'proyecto, no a ejecución bajo supervisión.', cursiva=True)

aviso('Una aclaración que evita malentendidos.',
      'El CITI publica también una tabla de honorarios mensuales por rol, de Bs 50 la hora. Esa '
      'es la referencia de un empleado en planilla, y no sirve para cotizar un trabajo por '
      'contrato: no cubre impuestos, herramientas, garantía ni el riesgo de entregar.')

doc.add_heading('4.2 Las horas', level=2)
p('El sistema terminado tiene 23 controladores, 27 modelos, 12 servicios de negocio, 69 vistas, '
  '98 rutas, 275 pruebas automáticas y una base de 29 tablas con 6 procedimientos y 7 '
  'disparadores: 33.695 líneas propias, sin contar librerías. Repartidas por bloque:')

tabla(['Bloque', 'Horas', 'Qué incluye'],
      [[nom, str(h), det] for nom, h, det in D.HORAS] +
      [['TOTAL', str(D.TOTAL_HORAS), 'Horas de desarrollo profesional acumuladas en el producto.']],
      anchos=[3.9, 1.4, 10.2], destacar=(len(D.HORAS),), derecha=(1,))

doc.add_heading('4.3 Lo que vale el sistema', level=2)
p('A la tarifa del arancel, el trabajo acumulado en este producto vale:')

tabla(['Concepto', 'Cálculo', 'US$', 'Bs'],
      [['Valor a arancel de técnico superior', '%d h × Bs %d' % (D.TOTAL_HORAS, D.TARIFA_BS),
        U(D.usd(D.VALOR_ARANCEL_BS)), B(D.VALOR_ARANCEL_BS)],
       ['Valor a arancel de técnico medio', '%d h × Bs 70' % D.TOTAL_HORAS,
        U(D.TOTAL_HORAS * 70 / D.TC), B(D.TOTAL_HORAS * 70)]],
      anchos=[6.0, 3.3, 3.1, 3.1], destacar=(0,), derecha=(2, 3))

aviso('Y sin embargo el precio no es ese. Aquí está la razón.',
      'Esa cifra es lo que costaría mandar a construir este sistema desde cero para un solo '
      'negocio. Pero el sistema YA EXISTE: el minimarket no paga el desarrollo, paga un producto '
      'terminado. Ese costo se reparte entre todos los negocios que lo usen, y lo que sí es '
      'exclusivo de esta instalación —ponerlo a andar, cargar el catálogo, capacitar— se cobra '
      'aparte. De ahí sale el precio de las páginas siguientes, muy por debajo del valor de '
      'arancel.', color='DCFCE7')

salto()

# ============================================================ 5. OPCIÓN A
doc.add_heading('5. Opción A — Compra del sistema completo', level=1)

p('El negocio compra el sistema, se instala en su servidor o en su computadora, y es suyo. Sin '
  'mensualidad obligatoria.', negrita=True)

precio(D.COMPRA_LANZAMIENTO_USD)

p('Precio de lanzamiento para el primer cliente. Precio de lista: %s.'
  % D.ambos(D.COMPRA_LISTA_USD), centro=True, tam=9.5, color=GRIS)

doc.add_heading('5.1 De dónde sale ese precio', level=2)
tabla(['Componente', 'Cálculo', 'US$', 'Bs'],
      [['Parte del desarrollo que le toca a este negocio',
        '%s ÷ %d negocios' % (D.b(D.VALOR_ARANCEL_BS), D.CLIENTES_AMORTIZACION),
        U(D.usd(D.AMORTIZACION_BS)), B(D.AMORTIZACION_BS)],
       ['Puesta en marcha de esta instalación',
        '%d h × Bs %d' % (D.HORAS_PUESTA, D.TARIFA_BS),
        U(D.usd(D.PUESTA_BS)), B(D.PUESTA_BS)],
       ['Costo del proyecto para este negocio', '',
        U(D.usd(D.COSTO_BASE_BS)), B(D.COSTO_BASE_BS)],
       ['PRECIO DE LANZAMIENTO', 'redondeado',
        U(D.COMPRA_LANZAMIENTO_USD), B(D.bs(D.COMPRA_LANZAMIENTO_USD))]],
      anchos=[6.0, 3.7, 2.9, 2.9], destacar=(2, 3), derecha=(2, 3))

p('Es el %s de lo que vale el sistema a arancel. Esa diferencia no es un descuento de buena '
  'voluntad: es la parte del desarrollo que pagan los demás negocios que también lo usen.'
  % D.pct(D.bs(D.COMPRA_LANZAMIENTO_USD) / D.VALOR_ARANCEL_BS * 100))

doc.add_heading('5.2 Las horas de la puesta en marcha', level=2)
tabla(['Tarea', 'Horas'],
      [[nom, str(h)] for nom, h in D.PUESTA_EN_MARCHA] +
      [['Total', str(D.HORAS_PUESTA)]],
      anchos=[12.5, 3.0], destacar=(len(D.PUESTA_EN_MARCHA),), derecha=(1,))

doc.add_heading('5.3 Qué se lleva el negocio', level=2)
vineta('para usarlo en el local sin límite de tiempo, de cajas ni de usuarios.',
       'Licencia perpetua: ')
vineta('el sistema completo, para no depender de nadie para seguir usándolo.',
       'Código fuente: ')
vineta('con el catálogo cargado, las cuentas creadas y el personal capacitado.',
       'Instalado y andando: ')
vineta('tres meses de correcciones sin costo desde la entrega.', 'Garantía: ')

doc.add_heading('5.4 Soporte después de los tres meses (opcional)', level=2)
p('Terminada la garantía, el negocio puede seguir solo o contratar soporte. No es obligatorio: el '
  'sistema es suyo y sigue funcionando igual.')
tabla(['Modalidad', 'US$', 'Bs', 'Qué cubre'],
      [['Soporte mensual', U(D.SOPORTE_OPCIONAL_USD), B(D.bs(D.SOPORTE_OPCIONAL_USD)),
        'Atención a consultas, correcciones y actualizaciones.'],
       ['Soporte anual', U(D.SOPORTE_OPCIONAL_USD * 10), B(D.bs(D.SOPORTE_OPCIONAL_USD * 10)),
        'Lo mismo, pagando diez meses y usando doce.']],
      anchos=[3.3, 2.3, 2.5, 7.4], derecha=(1, 2))

salto()

# ============================================================ 6. OPCIÓN B
doc.add_heading('6. Opción B — Suscripción de uso', level=1)

p('El negocio no compra nada: paga por usar el sistema, que queda alojado y mantenido por '
  'nosotros. Se entra desde el navegador y listo.', negrita=True)

precio(D.SUSCRIPCION_MES_USD, ' / mes')

p('Más una instalación inicial de %s, por única vez.' % D.ambos(D.INSTALACION_SUSC_USD),
  centro=True, tam=9.5, color=GRIS)

doc.add_heading('6.1 Las modalidades', level=2)
tabla(['Modalidad', 'US$', 'Bs', 'Detalle'],
      [['Instalación (una sola vez)', U(D.INSTALACION_SUSC_USD), B(D.bs(D.INSTALACION_SUSC_USD)),
        'Puesta en marcha, carga del catálogo y capacitación.'],
       ['Suscripción mensual', U(D.SUSCRIPCION_MES_USD), B(D.bs(D.SUSCRIPCION_MES_USD)),
        'Se paga por adelantado, mes a mes.'],
       ['Suscripción anual', U(D.SUSCRIPCION_ANIO_USD), B(D.bs(D.SUSCRIPCION_ANIO_USD)),
        'Paga diez meses y usa doce: %s por mes.' % D.ambos(D.SUSCRIPCION_ANIO_USD / 12.0)]],
      anchos=[3.9, 2.3, 2.5, 6.8], destacar=(2,), derecha=(1, 2))

doc.add_heading('6.2 Qué incluye la mensualidad', level=2)
vineta('servidor, dominio y certificado HTTPS. El negocio no compra ni administra nada.',
       'El alojamiento: ')
vineta('diarios y automáticos, guardados fuera del servidor. Si el disco falla, el negocio no '
       'pierde su historia.', 'Los respaldos: ')
vineta('cada mejora del sistema llega sin costo y sin trámite.', 'Las actualizaciones: ')
vineta('atención a consultas y corrección de errores durante toda la suscripción.', 'El soporte: ')
vineta('sin límite de tiempo, de cajas ni de usuarios, mientras la suscripción esté al día.',
       'El uso: ')

doc.add_heading('6.3 De dónde sale el precio mensual', level=2)
_infra = 7.0
_sop = D.TARIFA_BS / D.TC          # una hora de soporte al mes
_amort = D.usd(D.VALOR_ARANCEL_BS) / (D.CLIENTES_AMORTIZACION * D.MESES_PERMANENCIA)
_costo = _infra + _sop + _amort
tabla(['Componente', 'US$ / mes', 'Bs / mes'],
      [['Infraestructura compartida (servidor, dominio, respaldos)', U(_infra), B(D.bs(_infra))],
       ['Soporte y actualizaciones (una hora al mes)', U(_sop), B(D.bs(_sop))],
       ['Parte del desarrollo (%d negocios × %d meses)'
        % (D.CLIENTES_AMORTIZACION, D.MESES_PERMANENCIA), U(_amort), B(D.bs(_amort))],
       ['Costo', U(_costo), B(D.bs(_costo))],
       ['PRECIO', U(D.SUSCRIPCION_MES_USD), B(D.bs(D.SUSCRIPCION_MES_USD))]],
      anchos=[9.5, 3.0, 3.0], destacar=(3, 4), derecha=(1, 2))

aviso('Qué pasa si el negocio deja de pagar.',
      'El sistema deja de estar disponible, pero los datos no se borran: se entrega un respaldo '
      'completo de la base y de las fotos, en formato estándar, para que el negocio se lo lleve. '
      'La información del negocio es del negocio.')

salto()

# ============================================================ 7. COMPARACIÓN
doc.add_heading('7. Cuál de las dos conviene', level=1)

p('Depende de cuánto tiempo se piense usar el sistema, y de si el negocio quiere hacerse cargo de '
  'su propio servidor. Puesto en plata, a lo largo del tiempo:')

_equilibrio = (D.COMPRA_LANZAMIENTO_USD - D.INSTALACION_SUSC_USD) / float(D.SUSCRIPCION_MES_USD)
_filas = []
for meses in (6, 12, 24, 28, 36, 48, 60):
    _compra = D.COMPRA_LANZAMIENTO_USD
    _susc = D.INSTALACION_SUSC_USD + D.SUSCRIPCION_MES_USD * meses
    _filas.append(['%d meses' % meses, U(_compra), B(D.bs(_compra)), U(_susc), B(D.bs(_susc)),
                   'se igualan' if meses == 28 else
                   ('sale mejor suscribirse' if _susc < _compra else 'sale mejor comprar')])

tabla(['Al cabo de', 'Compra US$', 'Compra Bs', 'Susc. US$', 'Susc. Bs', ''],
      _filas, anchos=[2.1, 2.3, 2.5, 2.3, 2.5, 3.8], destacar=(3,), derecha=(1, 2, 3, 4))

p('El punto de equilibrio está en los %.0f meses: dos años y cuatro meses. Antes de eso la '
  'suscripción sale más barata; después, la compra.' % _equilibrio, negrita=True)

doc.add_heading('7.1 Puesto en criterios, no en plata', level=2)
tabla(['Si el negocio…', 'Le conviene'],
      [['Piensa usarlo por años y quiere que el sistema sea suyo', 'Comprar'],
       ['Prefiere no desembolsar mucho de entrada', 'Suscribirse'],
       ['No quiere hacerse cargo de un servidor ni de los respaldos', 'Suscribirse'],
       ['Ya tiene servidor propio, o quiere que los datos no salgan del local', 'Comprar'],
       ['Quiere probarlo de verdad antes de comprometerse', 'Suscribirse, y descontarlo si después compra'],
       ['Quiere el código fuente y no depender de nadie', 'Comprar']],
      anchos=[9.5, 6.0])

aviso('La suscripción se puede convertir en compra.',
      'Si el negocio arranca suscrito y más adelante decide comprar, se descuentan del precio de '
      'compra hasta doce meses de lo ya pagado. Probar primero no sale más caro.',
      color='DCFCE7')

salto()

# ============================================================ 8. CONDICIONES
doc.add_heading('8. Condiciones', level=1)

doc.add_heading('8.1 Forma de pago', level=2)
tabla(['Opción', 'Cómo se paga'],
      [['Compra', '50 % a la firma y 50 % contra entrega, una vez que el negocio revisó el sistema instalado y lo dio por bueno.'],
       ['Suscripción', 'La instalación a la firma. La mensualidad por adelantado, del 1 al 5 de cada mes. La anual, al inicio del período.']],
      anchos=[2.8, 12.7])

doc.add_heading('8.2 Plazos', level=2)
tabla(['Etapa', 'Plazo'],
      [['Instalación y configuración', '2 días hábiles desde la firma'],
       ['Carga del catálogo', '3 a 5 días hábiles desde que el negocio entrega su lista de productos'],
       ['Capacitación', '2 sesiones, dentro de la primera semana de uso'],
       ['Entrega formal', 'Al terminar la capacitación. Desde ahí corren los 3 meses de garantía.']],
      anchos=[4.8, 10.7])

doc.add_heading('8.3 Garantía y cambios', level=2)
vineta('durante tres meses desde la entrega, cualquier comportamiento del sistema que no '
       'corresponda a lo descrito en esta propuesta se corrige sin costo.', 'Garantía: ')
vineta('todo pedido que agregue alcance —un módulo nuevo, una integración, un reporte que hoy no '
       'existe— se cotiza aparte antes de hacerse, a la tarifa de arancel de técnico superior '
       '(%s la hora). Nada se empieza sin aprobación por escrito.'
       % D.ambos(D.TARIFA_BS / D.TC), 'Órdenes de cambio: ')
vineta('los datos que el negocio cargue son suyos. En cualquier momento se le entrega un '
       'respaldo completo en formato estándar.', 'Los datos: ')

doc.add_heading('8.4 Validez', level=2)
p('Esta propuesta es válida por 30 días a partir del 9 de septiembre de 2026. El tipo de cambio '
  'aplicado es de Bs 7,00 por dólar; si al momento de la firma difiere de forma significativa, '
  'los montos en bolivianos se recalculan al tipo de cambio vigente.')

doc.add_heading('Para cerrar', level=2)
p('El sistema está terminado y se puede probar hoy mismo en la demo, con datos reales de cuatro '
  'meses. La invitación es a que el personal lo use unos días —el cajero cobrando, el almacenero '
  'moviendo stock, la gerencia mirando reportes— y recién después decidir cuál de las dos '
  'opciones conviene.')
p('Cualquier duda sobre el alcance, el precio o los plazos se conversa antes de firmar.')

doc.save(D.salida('Propuesta-y-Presupuesto-Sistema-de-Ventas.docx'))
print('propuesta lista')
